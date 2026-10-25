import io
import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.submission import StudentSubmission
from app.models.examination import Examination
from app.models.question import Question
from app.models.evaluation import ExtractedAnswer, EvaluationResult, ApprovedMark
from app.core.config import REPORTS_DIR

class ReportGeneratorService:
    def get_student_report_data(self, db: Session, submission_id: int) -> Dict[str, Any]:
        """Fetch all details for student scorecard and reports."""
        submission = db.query(StudentSubmission).filter(StudentSubmission.id == submission_id).first()
        if not submission:
            raise HTTPException(status_code=404, detail="Student submission not found")

        exam = db.query(Examination).filter(Examination.id == submission.examination_id).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Examination not found")

        questions = db.query(Question).filter(Question.examination_id == exam.id).order_by(Question.order_index).all()
        answers = db.query(ExtractedAnswer).filter(ExtractedAnswer.submission_id == submission.id).all()
        approved_marks = db.query(ApprovedMark).filter(ApprovedMark.submission_id == submission.id).all()

        ans_by_q = {a.question_id: a for a in answers}
        app_by_q = {m.question_id: m for m in approved_marks}

        question_items = []
        total_approved = 0.0

        for q in questions:
            ans = ans_by_q.get(q.id)
            eval_res = ans.evaluation_result if ans else None
            app_mark = app_by_q.get(q.id)

            sug_marks = eval_res.suggested_marks if eval_res else 0.0
            t_marks = app_mark.teacher_marks if app_mark else None
            status_label = "Approved" if app_mark else ("Evaluated" if eval_res else "Unanswered")

            awarded = t_marks if t_marks is not None else 0.0
            total_approved += awarded

            question_items.append({
                "question_number": q.question_number,
                "question_text": q.question_text,
                "max_marks": q.max_marks,
                "suggested_marks": sug_marks,
                "teacher_marks": t_marks,
                "approved_marks": awarded,
                "status": status_label,
                "confidence_level": eval_res.confidence_level if eval_res else "N/A",
                "review_notes": app_mark.review_notes if app_mark else None
            })

        percentage = round((total_approved / exam.total_marks) * 100, 2) if exam.total_marks > 0 else 0.0

        return {
            "submission_id": submission.id,
            "examination_id": exam.id,
            "exam_title": exam.title,
            "subject": exam.subject,
            "class_name": exam.class_name,
            "academic_year": exam.academic_year,
            "exam_date": exam.exam_date,
            "total_exam_marks": exam.total_marks,
            "student_id": submission.student_id,
            "student_name": submission.student_name,
            "student_roll_no": submission.student_roll_no or "N/A",
            "total_approved_marks": total_approved,
            "percentage": percentage,
            "is_finalized": submission.is_finalized,
            "finalized_at": submission.finalized_at,
            "generated_at": datetime.now(timezone.utc),
            "questions": question_items
        }

    def generate_pdf_report(self, report_data: Dict[str, Any]) -> bytes:
        """Create a clean, formal academic PDF scorecard using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ExamTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#0F172A'),
            alignment=1  # Centered
        )
        subtitle_style = ParagraphStyle(
            'SubTitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#475569'),
            alignment=1
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#1E293B')
        )
        header_cell_style = ParagraphStyle(
            'HeaderCell',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.white
        )

        story = []

        # 1. Header
        story.append(Paragraph("GRADEAI — AUTOMATED EXAM EVALUATION SYSTEM", subtitle_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"Official Examination Scorecard: {report_data['exam_title']}", title_style))
        story.append(Paragraph(f"Subject: {report_data['subject']} | Class: {report_data['class_name']} | Academic Year: {report_data['academic_year']}", subtitle_style))
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A8A'), spaceBefore=5, spaceAfter=15))

        # 2. Student & Exam Meta Info Box
        meta_table_data = [
            [
                Paragraph(f"<b>Student Name:</b> {report_data['student_name']}", body_style),
                Paragraph(f"<b>Student ID:</b> {report_data['student_id']}", body_style),
            ],
            [
                Paragraph(f"<b>Roll No:</b> {report_data['student_roll_no']}", body_style),
                Paragraph(f"<b>Exam Date:</b> {report_data['exam_date']}", body_style),
            ],
            [
                Paragraph(f"<b>Total Score:</b> <font color='#1E3A8A'><b>{report_data['total_approved_marks']} / {report_data['total_exam_marks']}</b></font>", body_style),
                Paragraph(f"<b>Percentage:</b> <b>{report_data['percentage']}%</b>", body_style),
            ],
            [
                Paragraph(f"<b>Evaluation Status:</b> {'FINALIZED (Verified by Teacher)' if report_data['is_finalized'] else 'PENDING APPROVAL'}", body_style),
                Paragraph(f"<b>Generated On:</b> {report_data['generated_at'].strftime('%Y-%m-%d %H:%M UTC')}", body_style),
            ]
        ]
        meta_table = Table(meta_table_data, colWidths=[260, 260])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#E2E8F0')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 15))

        # 3. Question-Wise Breakdown Table
        story.append(Paragraph("<b>Question-by-Question Evaluation Breakdown</b>", styles['Heading3']))
        story.append(Spacer(1, 6))

        table_rows = [
            [
                Paragraph("<b>Q#</b>", header_cell_style),
                Paragraph("<b>Question Details</b>", header_cell_style),
                Paragraph("<b>Max</b>", header_cell_style),
                Paragraph("<b>AI Suggested</b>", header_cell_style),
                Paragraph("<b>Approved</b>", header_cell_style),
                Paragraph("<b>Status</b>", header_cell_style),
                Paragraph("<b>Teacher Notes</b>", header_cell_style)
            ]
        ]

        for q in report_data['questions']:
            q_text_short = (q['question_text'][:60] + '...') if len(q['question_text']) > 60 else q['question_text']
            table_rows.append([
                Paragraph(f"<b>{q['question_number']}</b>", body_style),
                Paragraph(q_text_short, body_style),
                Paragraph(str(q['max_marks']), body_style),
                Paragraph(f"{q['suggested_marks']:.1f}", body_style),
                Paragraph(f"<b>{q['approved_marks']:.1f}</b>", body_style),
                Paragraph(q['status'], body_style),
                Paragraph(q['review_notes'] or "-", body_style)
            ])

        q_table = Table(table_rows, colWidths=[35, 175, 40, 65, 55, 65, 95])
        q_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('ALIGN', (2, 0), (4, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        story.append(q_table)
        story.append(Spacer(1, 20))

        # 4. Teacher Verification Seal & Disclaimer
        disclaimer_text = (
            "<b>Human Oversight Statement:</b> This report reflects grades officially verified and approved "
            "by the course instructor. Automated AI marks are recommendations designed to reduce administrative workload "
            "and are not considered final without certified teacher approval."
        )
        disclaimer_p = Paragraph(disclaimer_text, ParagraphStyle('Disc', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor('#64748B')))
        story.append(disclaimer_p)

        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    def generate_csv_report(self, db: Session, exam_id: int) -> str:
        """Generate CSV export of all student marks for an exam."""
        exam = db.query(Examination).filter(Examination.id == exam_id).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Examination not found")

        submissions = db.query(StudentSubmission).filter(StudentSubmission.examination_id == exam_id).all()
        questions = db.query(Question).filter(Question.examination_id == exam_id).order_by(Question.order_index).all()

        output = io.StringIO()
        writer = csv.writer(output)

        # Header: Student Info + each Question + Total + Status
        headers = ["Student ID", "Student Name", "Roll No"]
        for q in questions:
            headers.append(f"Q{q.question_number} (Max {q.max_marks})")
        headers.extend(["Total Approved", "Percentage (%)", "Finalized Status"])
        writer.writerow(headers)

        for sub in submissions:
            row = [sub.student_id, sub.student_name, sub.student_roll_no or ""]
            approved_marks = db.query(ApprovedMark).filter(ApprovedMark.submission_id == sub.id).all()
            mark_dict = {m.question_id: m.teacher_marks for m in approved_marks}

            for q in questions:
                row.append(mark_dict.get(q.id, "N/A"))

            row.extend([sub.total_approved_marks, f"{sub.percentage:.2f}", "Finalized" if sub.is_finalized else "Draft"])
            writer.writerow(row)

        return output.getvalue()

    def generate_excel_report(self, db: Session, exam_id: int) -> bytes:
        """Generate formatted Excel report (.xlsx) with styles."""
        exam = db.query(Examination).filter(Examination.id == exam_id).first()
        if not exam:
            raise HTTPException(status_code=404, detail="Examination not found")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Exam Marks"

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title
        ws.append([f"GradeAI Assessment Report — {exam.title} ({exam.subject})"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Academic Year: {exam.academic_year} | Class: {exam.class_name} | Total Marks: {exam.total_marks}"])
        ws.append([])

        questions = db.query(Question).filter(Question.examination_id == exam_id).order_by(Question.order_index).all()
        headers = ["Student ID", "Student Name", "Roll No"]
        for q in questions:
            headers.append(f"Q{q.question_number} (Max {q.max_marks})")
        headers.extend(["Total Approved", "Percentage (%)", "Status"])

        ws.append(headers)
        header_row_idx = 4
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        submissions = db.query(StudentSubmission).filter(StudentSubmission.examination_id == exam_id).all()
        row_idx = 5
        for sub in submissions:
            approved_marks = db.query(ApprovedMark).filter(ApprovedMark.submission_id == sub.id).all()
            mark_dict = {m.question_id: m.teacher_marks for m in approved_marks}

            row_data = [sub.student_id, sub.student_name, sub.student_roll_no or ""]
            for q in questions:
                row_data.append(mark_dict.get(q.id, 0.0))
            row_data.extend([sub.total_approved_marks, sub.percentage, "Finalized" if sub.is_finalized else "Pending"])

            ws.append(row_data)
            for c_idx in range(1, len(row_data) + 1):
                ws.cell(row=row_idx, column=c_idx).border = thin_border
            row_idx += 1

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes


report_generator = ReportGeneratorService()
